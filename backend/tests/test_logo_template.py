"""Logo support in the DOCX quotation workflow.

Covers embedding the business logo into the DOCX header, saved position and
size settings, preserving the logo while replacing placeholders, SVG
rasterization, the show-logo toggle, and the no-logo fallback (business name
only).
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

from app.services.docx_templates import fill_docx_template
from tests.conftest import create_user, login, make_business_profile, quote_payload

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ── Helpers ────────────────────────────────────────────────────────────
def _svg_bytes() -> bytes:
    return (
        b'<svg xmlns="http://www.w3.org/2000/svg" width="40" height="24">'
        b'<rect width="40" height="24" fill="#2563eb"/></svg>'
    )


def _png_bytes() -> bytes:
    """A real, well-formed PNG rasterized from _svg_bytes().

    Hand-rolled minimal PNGs are tempting but easy to get wrong — python-docx
    walks PNG chunks strictly by their declared length fields, so a malformed
    fixture blows up in add_picture with a confusing UnicodeDecodeError.
    """
    from app.services.logo import svg_to_png

    png = svg_to_png(_svg_bytes())
    assert png is not None, "SVG rasterization should succeed in the test env"
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    return png


def _make_docx(paragraphs: list[str] | None = None, header: str | None = None) -> bytes:
    doc = Document()
    for p in paragraphs or []:
        doc.add_paragraph(p)
    if header is not None:
        doc.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _header_drawing(header):
    """Return (paragraph, drawing) for the first header image, else (None, None)."""
    for paragraph in header.paragraphs:
        for run in paragraph.runs:
            drawings = run._element.findall(qn("w:drawing"))
            if drawings:
                return paragraph, drawings[0]
    return None, None


def _drawing_height_emu(drawing) -> int:
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    return int(extent.get("cy"))


def _minimal_pdf() -> bytes:
    import reportlab.pdfgen.canvas as rl_canvas
    from reportlab.lib.pagesizes import A4

    out = io.BytesIO()
    c = rl_canvas.Canvas(out, pagesize=A4)
    c.drawString(50, 780, "FILLED DOCX RENDERED")
    c.showPage()
    c.save()
    return out.getvalue()


# ── Unit: header insertion + placeholder preservation ─────────────────
def test_logo_inserted_into_header():
    docx = _make_docx(["Body"], header="{{business_name}}")
    filled = fill_docx_template(
        docx, {"business_name": "Sparkle Cleaning"}, logo_bytes=_png_bytes()
    )
    doc = Document(io.BytesIO(filled))
    header = doc.sections[0].header
    para, drawing = _header_drawing(header)
    assert drawing is not None
    assert para.text == ""
    # Placeholder still replaced below the logo.
    texts = [p.text for p in header.paragraphs]
    assert "Sparkle Cleaning" in texts


def test_logo_preserved_while_placeholders_replaced():
    docx = _make_docx(["{{quote_number}}"], header="{{business_name}} — {{quote_number}}")
    filled = fill_docx_template(
        docx,
        {"business_name": "Sparkle", "quote_number": "Q-9"},
        logo_bytes=_png_bytes(),
    )
    doc = Document(io.BytesIO(filled))
    header = doc.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is not None
    assert any("Sparkle — Q-9" == p.text for p in header.paragraphs)
    assert doc.paragraphs[0].text == "Q-9"


@pytest.mark.parametrize(
    ("position", "expected"),
    [
        ("left", WD_ALIGN_PARAGRAPH.LEFT),
        ("center", WD_ALIGN_PARAGRAPH.CENTER),
        ("right", WD_ALIGN_PARAGRAPH.RIGHT),
    ],
)
def test_logo_position(position, expected):
    docx = _make_docx(header="{{business_name}}")
    filled = fill_docx_template(
        docx,
        {"business_name": "Sparkle"},
        logo_bytes=_png_bytes(),
        logo_position=position,
    )
    doc = Document(io.BytesIO(filled))
    para, drawing = _header_drawing(doc.sections[0].header)
    assert drawing is not None
    assert para.alignment == expected


@pytest.mark.parametrize(
    ("size", "inches"),
    [("small", 0.35), ("medium", 0.6), ("large", 0.9)],
)
def test_logo_size(size, inches):
    docx = _make_docx(header="{{business_name}}")
    filled = fill_docx_template(
        docx,
        {"business_name": "Sparkle"},
        logo_bytes=_png_bytes(),
        logo_size=size,
    )
    doc = Document(io.BytesIO(filled))
    _para, drawing = _header_drawing(doc.sections[0].header)
    assert _drawing_height_emu(drawing) == int(Inches(inches))


def test_no_logo_shows_business_name_only():
    docx = _make_docx(header="{{business_name}}")
    filled = fill_docx_template(docx, {"business_name": "Sparkle Cleaning"})
    doc = Document(io.BytesIO(filled))
    header = doc.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is None
    assert any(p.text == "Sparkle Cleaning" for p in header.paragraphs)


def test_show_logo_false_skips_logo():
    docx = _make_docx(header="{{business_name}}")
    filled = fill_docx_template(
        docx,
        {"business_name": "Sparkle Cleaning"},
        logo_bytes=_png_bytes(),
        show_logo=False,
    )
    doc = Document(io.BytesIO(filled))
    header = doc.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is None
    assert any(p.text == "Sparkle Cleaning" for p in header.paragraphs)


# ── Unit: SVG rasterization ────────────────────────────────────────────
def test_svg_to_png_returns_png():
    from app.services.logo import svg_to_png

    png = svg_to_png(_svg_bytes())
    assert png is not None
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_svg_to_png_invalid_returns_none():
    from app.services.logo import svg_to_png

    assert svg_to_png(b"not an svg at all") is None


def test_logo_bytes_for_docx_passthrough_for_png():
    from app.services.logo import logo_bytes_for_docx

    assert logo_bytes_for_docx(_png_bytes(), "image/png") == _png_bytes()


# ── Integration: PDF generation with logo ──────────────────────────────
async def _profile_with_docx(user_id: int, **overrides):
    from app.db.session import async_session_factory
    from app.models.business_profile import TemplateFile
    from app.uploads.store import store_template

    async with async_session_factory() as db:
        profile = make_business_profile(user_id, **overrides)
        db.add(profile)
        await db.flush()
        docx_bytes = _make_docx(["{{quote_number}}"], header="{{business_name}}")
        stored = store_template(DOCX_MIME, docx_bytes)
        tf = TemplateFile(
            user_id=user_id,
            stored_name=stored.name,
            original_name="template.docx",
            content_type=DOCX_MIME,
            size_bytes=len(docx_bytes),
        )
        db.add(tf)
        await db.flush()
        profile.template_type = "custom_docx"
        profile.custom_docx_file_id = tf.id
        profile.custom_docx_filename = "template.docx"
        profile.custom_docx_mime_type = DOCX_MIME
        profile.custom_docx_size = tf.size_bytes
        profile.custom_docx_uploaded_at = datetime.now(UTC)
        profile.custom_docx_version = 1
        await db.commit()
        await db.refresh(profile)
        return profile


def _capture_converter(monkeypatch):
    from app.services import docx_to_pdf

    captured: dict = {}

    def fake_convert(docx_bytes: bytes) -> bytes:
        captured["docx"] = docx_bytes
        return _minimal_pdf()

    monkeypatch.setattr(docx_to_pdf, "is_converter_available", lambda: True)
    monkeypatch.setattr(docx_to_pdf, "convert_docx_to_pdf", fake_convert)
    return captured


async def _create_quote(client, **overrides):
    res = await client.post("/api/quotes", json=quote_payload(**overrides))
    assert res.status_code == 201, res.text
    return res.json()["quote"]


async def _upload_logo(client, data: bytes, filename: str, mime: str):
    files = {"file": (filename, io.BytesIO(data), mime)}
    return await client.post("/api/business-profile/logo", files=files)


@pytest.mark.asyncio
async def test_generate_pdf_embeds_uploaded_logo(client, monkeypatch):
    captured = _capture_converter(monkeypatch)
    user = await create_user("logo-pdf@example.com")
    await login(client, "logo-pdf@example.com")
    await _profile_with_docx(user.id, tax_rate_bps=750)

    assert (await _upload_logo(client, _png_bytes(), "logo.png", "image/png")).status_code == 200

    quote = await _create_quote(client, quote_number="Q-LOGO")
    res = await client.post(f"/api/quotes/{quote['id']}/generate-pdf")
    assert res.status_code == 200, res.text
    assert res.headers.get("x-quoteflow-template") == "custom_docx"

    filled = Document(io.BytesIO(captured["docx"]))
    header = filled.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is not None
    assert any("Sparkle Cleaning" in p.text for p in header.paragraphs)


@pytest.mark.asyncio
async def test_generate_pdf_svg_logo_rasterized(client, monkeypatch):
    captured = _capture_converter(monkeypatch)
    user = await create_user("logo-svgpdf@example.com")
    await login(client, "logo-svgpdf@example.com")
    await _profile_with_docx(user.id, tax_rate_bps=750)

    res = await _upload_logo(client, _svg_bytes(), "logo.svg", "image/svg+xml")
    assert res.status_code == 200, res.text

    quote = await _create_quote(client, quote_number="Q-SVGLOGO")
    res = await client.post(f"/api/quotes/{quote['id']}/generate-pdf")
    assert res.status_code == 200, res.text
    assert res.headers.get("x-quoteflow-template") == "custom_docx"

    filled = Document(io.BytesIO(captured["docx"]))
    _para, drawing = _header_drawing(filled.sections[0].header)
    assert drawing is not None


@pytest.mark.asyncio
async def test_generate_pdf_toggle_off_omits_logo(client, monkeypatch):
    captured = _capture_converter(monkeypatch)
    user = await create_user("logo-off@example.com")
    await login(client, "logo-off@example.com")
    await _profile_with_docx(user.id, tax_rate_bps=750)
    assert (await _upload_logo(client, _png_bytes(), "logo.png", "image/png")).status_code == 200

    updated = await client.put("/api/business-profile", json={"show_logo_on_quotation": False})
    assert updated.status_code == 200
    assert updated.json()["show_logo_on_quotation"] is False

    quote = await _create_quote(client, quote_number="Q-LOGOOFF")
    res = await client.post(f"/api/quotes/{quote['id']}/generate-pdf")
    assert res.status_code == 200, res.text

    filled = Document(io.BytesIO(captured["docx"]))
    header = filled.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is None
    assert any("Sparkle Cleaning" in p.text for p in header.paragraphs)


@pytest.mark.asyncio
async def test_generate_pdf_no_logo_falls_back_to_business_name(client, monkeypatch):
    captured = _capture_converter(monkeypatch)
    user = await create_user("logo-none@example.com")
    await login(client, "logo-none@example.com")
    await _profile_with_docx(user.id, tax_rate_bps=750)

    quote = await _create_quote(client, quote_number="Q-NOLOGO")
    res = await client.post(f"/api/quotes/{quote['id']}/generate-pdf")
    assert res.status_code == 200, res.text

    filled = Document(io.BytesIO(captured["docx"]))
    header = filled.sections[0].header
    _para, drawing = _header_drawing(header)
    assert drawing is None
    assert any("Sparkle Cleaning" in p.text for p in header.paragraphs)


@pytest.mark.asyncio
async def test_generate_pdf_logo_position_and_size_settings_applied(client, monkeypatch):
    captured = _capture_converter(monkeypatch)
    user = await create_user("logo-pos@example.com")
    await login(client, "logo-pos@example.com")
    await _profile_with_docx(user.id, tax_rate_bps=750)
    assert (await _upload_logo(client, _png_bytes(), "logo.png", "image/png")).status_code == 200

    updated = await client.put(
        "/api/business-profile", json={"logo_position": "right", "logo_size": "small"}
    )
    assert updated.status_code == 200

    quote = await _create_quote(client, quote_number="Q-LOGOPOS")
    res = await client.post(f"/api/quotes/{quote['id']}/generate-pdf")
    assert res.status_code == 200, res.text

    filled = Document(io.BytesIO(captured["docx"]))
    para, drawing = _header_drawing(filled.sections[0].header)
    assert drawing is not None
    assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert _drawing_height_emu(drawing) == int(Inches(0.35))
