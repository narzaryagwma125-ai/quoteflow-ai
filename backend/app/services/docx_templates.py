"""DOCX template filling.

Placeholders (e.g. ``{{business_name}}``) are replaced inside paragraphs,
table cells, headers and footers without destroying formatting where possible.
The special ``{{line_items}}`` placeholder is replaced by a Word table.
"""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from app.services.currency import currency_symbol as _shared_currency_symbol

PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")

_LOGO_POSITION_ALIGN: dict[str, int] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}
_LOGO_SIZES_INCHES: dict[str, float] = {"small": 0.35, "medium": 0.6, "large": 0.9}


def currency_symbol(currency: str) -> str:
    """Currency symbol prefix, falling back to the ISO code."""
    return _shared_currency_symbol(currency)


def format_minor(minor: int) -> str:
    """Minor units (cents) → '1,234.56'."""
    return f"{Decimal(minor) / 100:,.2f}"


def format_currency(minor: int, currency: str) -> str:
    return f"{currency_symbol(currency)}{format_minor(minor)}"


def format_date(value: str) -> str:
    """ISO date → 'September 13, 2026'. Returns input untouched on parse failure."""
    if not value:
        return ""
    try:
        parts = value.split("-")
        return date(int(parts[0]), int(parts[1]), int(parts[2])).strftime("%B %d, %Y")
    except (ValueError, IndexError):
        return value


def _replace_text(text: str, values: dict[str, str]) -> str:
    def _rep(m: re.Match) -> str:
        key = m.group(1)
        if key == "line_items":
            # Keep the placeholder intact — handled separately as a table.
            return m.group(0)
        return values.get(key, "")

    return PLACEHOLDER_RE.sub(_rep, text)


def _replace_paragraph(paragraph, values: dict[str, str]) -> None:
    """Replace placeholders in a paragraph, preserving run formatting."""
    if not PLACEHOLDER_RE.search(paragraph.text):
        return

    new_text = _replace_text(paragraph.text, values)
    if not paragraph.runs:
        paragraph.text = new_text
        return

    replaced = False
    for run in paragraph.runs:
        if PLACEHOLDER_RE.search(run.text):
            run.text = _replace_text(run.text, values)
            replaced = True

    if not replaced:
        # A placeholder was split across runs — merge into the first run.
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""


def _table_paragraphs(tables) -> list:
    """All paragraphs from a list of tables, recursing into nested tables."""
    out = []
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
                if cell.tables:
                    out.extend(_table_paragraphs(cell.tables))
    return out


def _iter_text_containers(doc: Document):
    """Yield body, table-cell, header and footer paragraphs."""
    yield from doc.paragraphs
    yield from _table_paragraphs(doc.tables)

    for section in doc.sections:
        containers = [section.header, section.footer]
        try:
            containers += [
                section.first_page_header,
                section.first_page_footer,
                section.even_page_header,
                section.even_page_footer,
            ]
        except AttributeError:
            pass
        for container in containers:
            if container is None:
                continue
            yield from container.paragraphs
            yield from _table_paragraphs(container.tables)


def _add_logo_to_header(header, logo_bytes: bytes, position: str, size: str) -> None:
    """Insert the logo as a new image paragraph at the top of a header."""
    paragraphs = header.paragraphs
    target = paragraphs[0] if paragraphs else header.add_paragraph()
    if target.text.strip():
        target = target.insert_paragraph_before()
    run = target.add_run()
    run.add_picture(io.BytesIO(logo_bytes), height=Inches(_LOGO_SIZES_INCHES.get(size, 0.6)))
    target.alignment = _LOGO_POSITION_ALIGN.get(position, WD_ALIGN_PARAGRAPH.LEFT)


def _insert_logo_in_header(doc: Document, logo_bytes: bytes, position: str, size: str) -> None:
    """Insert the logo image into every applicable header in the document.

    Headers are resolved to their underlying part and de-duplicated so an
    inherited (linked) header is not filled more than once.
    """
    seen: set[str] = set()
    for section in doc.sections:
        targets = [section.header]
        if section.different_first_page_header_footer:
            targets.append(section.first_page_header)
        if getattr(doc.settings, "odd_and_even_pages_header_footer", False):
            targets.append(section.odd_page_header)
            targets.append(section.even_page_header)
        for header in targets:
            partname = str(header.part.partname)
            if partname in seen:
                continue
            seen.add(partname)
            _add_logo_to_header(header, logo_bytes, position, size)


def _build_line_items_table(doc: Document, items: list[dict], currency: str) -> bool:
    """Replace the ``{{line_items}}`` paragraph with a Word table. Returns True
    when a table was inserted."""
    for paragraph in list(doc.paragraphs):
        if "{{line_items}}" not in paragraph.text:
            continue

        table = doc.add_table(rows=1, cols=6)
        try:
            table.style = "Table Grid"
        except Exception:
            pass  # Style not defined in the template — fall back to plain.

        headers = ["#", "Description", "Quantity", "Unit", "Unit Price", "Amount"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for idx, item in enumerate(items, start=1):
            row = table.add_row()
            row.cells[0].text = str(idx)
            row.cells[1].text = str(item.get("description", ""))
            row.cells[2].text = str(item.get("quantity", ""))
            row.cells[3].text = str(item.get("unit", ""))
            row.cells[4].text = format_currency(int(item.get("unit_price_minor", 0)), currency)
            row.cells[5].text = format_currency(int(item.get("line_total_minor", 0)), currency)

        paragraph._p.addnext(table._tbl)
        paragraph._p.getparent().remove(paragraph._p)
        return True
    return False


def fill_docx_template(
    docx_bytes: bytes,
    values: dict[str, str],
    items: list[dict] | None = None,
    logo_bytes: bytes | None = None,
    logo_position: str = "left",
    logo_size: str = "medium",
    show_logo: bool = True,
) -> bytes:
    """Fill a DOCX template in memory and return the filled document bytes."""
    doc = Document(io.BytesIO(docx_bytes))

    # Insert the logo *before* placeholder replacement so the image run is
    # never merged or overwritten by text-merging logic.
    if logo_bytes and show_logo:
        _insert_logo_in_header(doc, logo_bytes, logo_position, logo_size)

    for paragraph in _iter_text_containers(doc):
        _replace_paragraph(paragraph, values)

    if items:
        _build_line_items_table(doc, items, values.get("currency", "USD"))

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_sample_docx() -> bytes:
    """Create a sample DOCX template containing every supported placeholder."""
    doc = Document()

    doc.add_heading("Quotation", 0)

    doc.add_heading("Business Information", level=2)
    doc.add_paragraph("{{business_name}}")
    doc.add_paragraph("{{business_owner}}")
    doc.add_paragraph("{{business_email}} | {{business_phone}}")
    doc.add_paragraph("{{business_address}}")
    doc.add_paragraph("Tax ID: {{business_tax_id}}")
    doc.add_paragraph("Currency: {{business_currency}}")

    doc.add_heading("Quote Details", level=2)
    doc.add_paragraph("Quote #: {{quote_number}}")
    doc.add_paragraph("Issue date: {{issue_date}}")
    doc.add_paragraph("Valid until: {{valid_until}}")
    doc.add_paragraph("Status: {{status}}")
    doc.add_paragraph("Currency: {{currency}}")

    doc.add_heading("Customer", level=2)
    doc.add_paragraph("{{customer_name}}")
    doc.add_paragraph("{{customer_company}}")
    doc.add_paragraph("{{customer_email}}")
    doc.add_paragraph("{{customer_phone}}")
    doc.add_paragraph("{{customer_address}}")

    doc.add_heading("Line Items", level=2)
    doc.add_paragraph("{{line_items}}")

    doc.add_heading("Totals", level=2)
    doc.add_paragraph("Subtotal: {{subtotal}}")
    doc.add_paragraph("Discount: {{discount}}")
    doc.add_paragraph("Tax rate: {{tax_rate}}")
    doc.add_paragraph("Tax amount: {{tax_amount}}")
    doc.add_paragraph("TOTAL: {{total}}")

    doc.add_heading("Additional Information", level=2)
    doc.add_paragraph("Notes: {{notes}}")
    doc.add_paragraph("Payment terms: {{payment_terms}}")
    doc.add_paragraph("Signature: {{signature}}")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
